# receipt_exporter.py
"""
单据生成器
将出入库明细生成A5横向的Excel单据，支持分组分页
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


class ReceiptExporter:
    # 列宽配置（单位：字符）
    COL_WIDTHS = {
        'A': 5.5, 'B': 9, 'C': 10.5, 'D': 10.5, 'E': 12,
        'F': 6.5, 'G': 6, 'H': 9, 'I': 9, 'J': 12
    }

    @staticmethod
    def _create_styles():
        """创建所有要用到的样式对象，只创建一次"""
        return {
            'title_font': Font(name='宋体', size=14, bold=True),
            'title_align': Alignment(horizontal='center', vertical='center'),
            'info_font': Font(name='宋体', size=10),
            'info_align_left': Alignment(horizontal='left', vertical='center'),
            'info_align_right': Alignment(horizontal='right', vertical='center'),
            'header_font': Font(name='宋体', size=9, bold=True),
            'header_align': Alignment(horizontal='center', vertical='center', wrap_text=True),
            'header_border': Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            ),
            'header_fill': PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid'),
            'data_font': Font(name='宋体', size=9),
            'data_align': Alignment(horizontal='center', vertical='center'),
            'data_border': Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            ),
            'sign_label_font': Font(name='宋体', size=10, bold=True),
            'sign_label_align': Alignment(horizontal='left', vertical='center'),
            'sign_line_font': Font(name='宋体', size=10, underline='single'),
        }

    @classmethod
    def _setup_a5_sheet(cls, ws):
        """设置A5横向页面参数"""
        # A5纸张，横向
        ws.page_setup.paperSize = 11
        ws.page_setup.orientation = 'landscape'

        # 页边距（英寸）
        ws.page_margins.left = 0.5
        ws.page_margins.right = 0.5
        ws.page_margins.top = 0.6
        ws.page_margins.bottom = 0.6

        # 列宽
        for col_letter, width in cls.COL_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width

        # 标题行和表头行高
        for r, h in enumerate([28, 22, 22, 22], 1):
            ws.row_dimensions[r].height = h

    @classmethod
    def _write_data_rows(cls, ws, styles, page_data, start_no):
        """写入数据行（从第5行开始，最多10行）"""
        for row_offset, item in enumerate(page_data):
            row = 5 + row_offset
            values = [
                start_no + row_offset,
                item.get('belong_equipment', ''),
                item.get('name', ''),
                item.get('spec_model', ''),
                item.get('serial_no', ''),
                item.get('quantity', 0),
                item.get('unit', ''),
                round(item.get('unit_price', 0), 2),
                round(item.get('total_price', 0), 2),
                item.get('remark', '')
            ]
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=ci, value=val)
                cell.font = styles['data_font']
                cell.alignment = styles['data_align']
                cell.border = styles['data_border']
            ws.row_dimensions[row].height = 20

        # 不足10行时补空行（保持表格完整）
        for row_offset in range(10 - len(page_data)):
            row = 5 + len(page_data) + row_offset
            for ci in range(1, 11):
                cell = ws.cell(row=row, column=ci)
                cell.font = styles['data_font']
                cell.border = styles['data_border']
            ws.row_dimensions[row].height = 20

    @classmethod
    def _write_sign_row(cls, ws, styles, sign_row, left_label, left_value, right_label, right_value):
        """写入签名行"""
        ws.row_dimensions[sign_row].height = 30
        # 左侧签名区
        ws.merge_cells(f'A{sign_row}:E{sign_row}')
        left_text = f"{left_label}：{left_value}" if left_value else f"{left_label}：{'__'*6}"
        cell = ws.cell(row=sign_row, column=1, value=left_text)
        cell.font = styles['sign_label_font'] if left_value else styles['sign_line_font']
        cell.alignment = styles['sign_label_align']

        # 右侧签名区
        ws.merge_cells(f'G{sign_row}:J{sign_row}')
        right_text = f"{right_label}：{right_value}" if right_value else f"{right_label}：{'__'*6}"
        cell = ws.cell(row=sign_row, column=7, value=right_text)
        cell.font = styles['sign_label_font'] if right_value else styles['sign_line_font']
        cell.alignment = styles['sign_label_align']

    @classmethod
    def generate_in_receipt(cls, items, output_dir, group_name, contract_no, in_date,
                            stock_in_person='', operator='', group_type='factory'):
        """
        生成入库单Excel文件
        :param items: 入库明细列表，每个元素为包含器材信息的字典
        :param output_dir: 输出文件夹
        :param group_name: 分组名称（厂家名或日期）
        :param contract_no: 合同编号
        :param in_date: 入库日期字符串
        :param stock_in_person: 入库人
        :param operator: 经办人
        :param group_type: 分组方式 'factory' 或 'date'
        :return: 生成的文件路径列表
        """
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_name = (group_name or '未知').replace('/', '_').replace('\\', '_')
        prefix = "器材入库单"

        # 每页10行，分页
        page_items = [items[i:i + 10] for i in range(0, len(items), 10)]
        styles = cls._create_styles()
        filepaths = []

        for page_num, page_data in enumerate(page_items, 1):
            start_no = (page_num - 1) * 10 + 1
            total_pages = len(page_items)

            # 构造文件名
            filename = f"{prefix}_{safe_name}_{timestamp}"
            if total_pages > 1:
                filename += f"_第{page_num}页"
            filename += ".xlsx"
            filepath = os.path.join(output_dir, filename)

            wb = Workbook()
            ws = wb.active
            ws.title = "器材入库单"
            cls._setup_a5_sheet(ws)

            # ===== 第1行：标题 =====
            ws.merge_cells('A1:J1')
            c = ws.cell(row=1, column=1, value="器 材 入 库 单")
            c.font = styles['title_font']
            c.alignment = styles['title_align']

            # ===== 第2行：分组信息 =====
            if group_type == 'factory':
                left_label = "生产厂家："
                left_value = group_name
                right_label = "入库日期："
                right_value = in_date
            else:
                left_label = "入库日期："
                left_value = group_name
                right_label = "生产厂家："
                right_value = items[0].get('manufacturer', '')

            ws.merge_cells('A2:E2')
            cell = ws.cell(row=2, column=1, value=f"{left_label}{left_value}")
            cell.font = styles['info_font']
            cell.alignment = styles['info_align_left']

            ws.merge_cells('G2:H2')
            cell = ws.cell(row=2, column=7, value=right_label)
            cell.font = styles['info_font']
            cell.alignment = styles['info_align_right']

            ws.merge_cells('I2:J2')
            cell = ws.cell(row=2, column=9, value=right_value)
            cell.font = styles['info_font']
            cell.alignment = styles['info_align_left']

            # ===== 第3行：合同编号 + 页码 =====
            ws.merge_cells('A3:E3')
            cell = ws.cell(row=3, column=1, value=f"合同编号：{contract_no}")
            cell.font = styles['info_font']
            cell.alignment = styles['info_align_left']

            if total_pages > 1:
                ws.merge_cells('G3:J3')
                cell = ws.cell(row=3, column=7, value=f"第 {page_num}/{total_pages} 页")
                cell.font = Font(name='宋体', size=9)
                cell.alignment = styles['info_align_right']

            # ===== 第4行：表头 =====
            headers = ['序号', '所属装备', '器材名称', '规格型号', '器材编号',
                       '数量', '计量单位', '单价（元）', '总价（元）', '备注']
            for ci, h in enumerate(headers, 1):
                cell = ws.cell(row=4, column=ci, value=h)
                cell.font = styles['header_font']
                cell.alignment = styles['header_align']
                cell.border = styles['header_border']
                cell.fill = styles['header_fill']

            # ===== 数据行 =====
            cls._write_data_rows(ws, styles, page_data, start_no)

            # ===== 签名行（第15行）=====
            cls._write_sign_row(ws, styles, 15, '入库人', stock_in_person, '经办人', operator)

            # 打印区域与缩放
            ws.print_area = 'A1:J15'
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1

            wb.save(filepath)
            wb.close()
            filepaths.append(filepath)

        return filepaths

    @classmethod
    def generate_out_receipt(cls, items, output_dir, group_name, contract_no, out_date,
                             stock_out_person='', operator='', group_type='equipment'):
        """
        生成出库单Excel文件
        :param items: 出库明细列表
        :param output_dir: 输出文件夹
        :param group_name: 分组名称（所属装备或日期）
        :param contract_no: 合同编号
        :param out_date: 出库日期字符串
        :param stock_out_person: 出库人
        :param operator: 经办人
        :param group_type: 分组方式 'equipment' 或 'date'
        :return: 文件路径列表
        """
        try:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            safe_name = (group_name or '未知').replace('/', '_').replace('\\', '_')
            page_items = [items[i:i + 10] for i in range(0, len(items), 10)]
            styles = cls._create_styles()
            filepaths = []

            for page_num, page_data in enumerate(page_items, 1):
                start_no = (page_num - 1) * 10 + 1
                total_pages = len(page_items)
                filename = f"器材出库单_{safe_name}_{timestamp}"
                if total_pages > 1:
                    filename += f"_第{page_num}页"
                filename += ".xlsx"
                filepath = os.path.join(output_dir, filename)

                wb = Workbook()
                ws = wb.active
                ws.title = "器材出库单"
                cls._setup_a5_sheet(ws)

                # 标题
                ws.merge_cells('A1:J1')
                c = ws.cell(row=1, column=1, value="器 材 出 库 单")
                c.font = styles['title_font']
                c.alignment = styles['title_align']

                # 分组信息
                if group_type == 'equipment':
                    left_label = "所属装备："
                    left_value = group_name
                    right_label = "出库日期："
                    right_value = out_date
                else:
                    left_label = "出库日期："
                    left_value = group_name
                    right_label = "所属装备："
                    right_value = items[0].get('belong_equipment', '')

                ws.merge_cells('A2:E2')
                cell = ws.cell(row=2, column=1, value=f"{left_label}{left_value}")
                cell.font = styles['info_font']
                cell.alignment = styles['info_align_left']

                ws.merge_cells('G2:H2')
                cell = ws.cell(row=2, column=7, value=right_label)
                cell.font = styles['info_font']
                cell.alignment = styles['info_align_right']

                ws.merge_cells('I2:J2')
                cell = ws.cell(row=2, column=9, value=right_value)
                cell.font = styles['info_font']
                cell.alignment = styles['info_align_left']

                # 合同编号
                ws.merge_cells('A3:E3')
                cell = ws.cell(row=3, column=1, value=f"合同编号：{contract_no}")
                cell.font = styles['info_font']
                cell.alignment = styles['info_align_left']

                if total_pages > 1:
                    ws.merge_cells('G3:J3')
                    cell = ws.cell(row=3, column=7, value=f"第 {page_num}/{total_pages} 页")
                    cell.font = Font(name='宋体', size=9)
                    cell.alignment = styles['info_align_right']

                # 表头
                headers = ['序号', '所属装备', '器材名称', '规格型号', '器材编号',
                           '数量', '计量单位', '单价（元）', '总价（元）', '备注']
                for ci, h in enumerate(headers, 1):
                    cell = ws.cell(row=4, column=ci, value=h)
                    cell.font = styles['header_font']
                    cell.alignment = styles['header_align']
                    cell.border = styles['header_border']
                    cell.fill = styles['header_fill']

                # 数据行
                cls._write_data_rows(ws, styles, page_data, start_no)

                # 签名行
                cls._write_sign_row(ws, styles, 15, '出库人', stock_out_person, '经办人', operator)

                ws.print_area = 'A1:J15'
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 1

                wb.save(filepath)
                wb.close()
                filepaths.append(filepath)

            return filepaths
        except Exception as e:
            print(f"生成出库单时发生错误: {e}")
            import traceback
            traceback.print_exc()
            raise